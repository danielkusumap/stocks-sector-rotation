"""Import `saham_sektoral.xlsx` → tabel `stock` (universe). Aturan parsing: docs/02 §6."""

from __future__ import annotations

import hashlib
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.workbook import Workbook

from .. import db

TICKER_OK = re.compile(r"[A-Z0-9]{4}\Z")
ID_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "mei": 5, "jun": 6,
    "jul": 7, "agt": 8, "sep": 9, "okt": 10, "nov": 11, "des": 12,
}
# Header Excel → nama kolom internal (dicocokkan case-insensitive), plus posisi cadangan
# bila header tidak dikenal: file nyata = No|Kode|Nama Perusahaan|Tanggal Pencatatan|Saham|Papan Pencatatan.
HEADER_ALIAS = {
    "kode": "ticker",
    "nama perusahaan": "name",
    "tanggal pencatatan": "listing_date",
    "saham": "shares_outstanding",
    "papan pencatatan": "board",
}
FALLBACK_POS = {"ticker": 1, "name": 2, "listing_date": 3, "shares_outstanding": 4, "board": 5}

ERROR_REASONS = ("invalid_ticker", "unknown_sheet", "duplicate_ticker", "empty_row")


class AlreadyImported(Exception):
    """File dengan sha256 sama sudah pernah diimport (HTTP 409) kecuali ?force=true."""

    def __init__(self, import_id: int, sha256: str):
        super().__init__(f"file sudah pernah diimport (import_id={import_id})")
        self.import_id = import_id
        self.sha256 = sha256


class UnknownSheet(Exception):
    """Nama sheet di luar 11 kode sektor; hanya fatal bila ?strict=true."""

    def __init__(self, sheet: str):
        super().__init__(f"sheet '{sheet}' bukan kode sektor IDX-IC")
        self.sheet = sheet


class NoValidRows(Exception):
    """File tidak menghasilkan satu saham pun — jangan dipakai untuk menonaktifkan universe."""


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def normalize_ticker(raw: object) -> str | None:
    """Uppercase, buang `.JK`/spasi/titik; valid hanya 4 karakter [A-Z0-9]."""
    if raw is None:
        return None
    s = re.sub(r"[^A-Z0-9]", "", str(raw).strip().upper().replace(".JK", ""))
    return s if TICKER_OK.match(s) else None


def parse_listing_date(raw: object) -> str | None:
    """`06 Jul 2021` (bulan Indonesia) → `2021-07-06`; gagal parse = None, bukan error."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    s = str(raw).strip()
    m = re.match(r"(\d{1,2})\s+([A-Za-z]{3})\w*\s+(\d{4})", s)
    if m:
        day, mon, year = m.groups()
        mi = ID_MONTHS.get(mon.lower())
        return f"{year}-{mi:02d}-{int(day):02d}" if mi else None
    return s if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s) else None


def parse_shares(raw: object) -> int | None:
    """Jumlah lembar. Terima angka Excel, `8.603.416.176` (format ID), atau string biasa."""
    if raw is None or str(raw).strip() == "":
        return None
    if isinstance(raw, (int, float)):
        return int(raw)
    s = str(raw).strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        return int(s.replace(".", ""))
    try:
        return int(float(s))
    except ValueError:
        return None


def _cell(row: tuple, idx: dict[str, int], name: str) -> object:
    i = idx.get(name, FALLBACK_POS[name]) if len(idx) else FALLBACK_POS[name]
    return row[i] if i < len(row) else None


def parse_workbook(
    wb: Workbook, sector_ids: dict[str, int], *, strict: bool = False
) -> tuple[list[dict], list[dict]]:
    """Baca semua sheet → (baris siap upsert, daftar error)."""
    rows: list[dict] = []
    errors: list[dict] = []
    seen: set[str] = set()

    for ws in wb.worksheets:
        sector_id = sector_ids.get(str(ws.title).strip().upper())
        if sector_id is None:
            if strict:
                raise UnknownSheet(ws.title)
            errors.append({"sheet": ws.title, "row_no": None, "raw_value": None, "reason": "unknown_sheet"})
            continue

        header = next(ws.iter_rows(values_only=True), ())
        idx = {
            HEADER_ALIAS[str(c).strip().lower()]: i
            for i, c in enumerate(header)
            if c is not None and str(c).strip().lower() in HEADER_ALIAS
        }

        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None or str(v).strip() == "" for v in row):
                continue  # baris kosong di ekor sheet = artefak Excel, bukan error

            raw_kode = _cell(row, idx, "ticker")
            ticker = normalize_ticker(raw_kode)
            if ticker is None:
                blank = raw_kode is None or str(raw_kode).strip() == ""
                errors.append(
                    {
                        "sheet": ws.title,
                        "row_no": row_no,
                        "raw_value": None if blank else str(raw_kode)[:50],
                        "reason": "empty_row" if blank else "invalid_ticker",
                    }
                )
                continue
            if ticker in seen:
                errors.append(
                    {"sheet": ws.title, "row_no": row_no, "raw_value": ticker, "reason": "duplicate_ticker"}
                )
                continue
            seen.add(ticker)

            name = _cell(row, idx, "name")
            board = _cell(row, idx, "board")
            rows.append(
                {
                    "ticker": ticker,
                    "sector_id": sector_id,
                    "name": str(name).strip() if name is not None and str(name).strip() else None,
                    "board": str(board).strip() if board is not None and str(board).strip() else None,
                    "listing_date": parse_listing_date(_cell(row, idx, "listing_date")),
                    "shares_outstanding": parse_shares(_cell(row, idx, "shares_outstanding")),
                }
            )
    return rows, errors


def import_excel(
    data: bytes,
    filename: str,
    *,
    force: bool = False,
    strict: bool = False,
    db_path: str | Path | None = None,
) -> dict:
    """Import idempotent per sha256. Saham yang hilang dari file → `is_active = 0` (tidak dihapus)."""
    sha = hashlib.sha256(data).hexdigest()
    conn = db.connect(db_path)
    try:
        prev = conn.execute("SELECT id FROM excel_import WHERE sha256 = ?", (sha,)).fetchone()
        if prev and not force:
            raise AlreadyImported(prev["id"], sha)

        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            sheet_count = len(wb.sheetnames)
            sector_ids = {r["code"]: r["id"] for r in conn.execute("SELECT id, code FROM sector")}
            rows, errors = parse_workbook(wb, sector_ids, strict=strict)
        finally:
            wb.close()

        if not rows:
            raise NoValidRows("tidak ada baris saham valid di file ini")

        ts = now_iso()
        inserted = updated = 0
        for r in rows:
            exists = conn.execute("SELECT 1 FROM stock WHERE ticker = ?", (r["ticker"],)).fetchone()
            if exists:
                updated += 1
            else:
                inserted += 1
            # shares_outstanding & listing_date ikut di-UPDATE setiap import (aksi korporasi).
            conn.execute(
                """
                INSERT INTO stock (ticker, yahoo_symbol, name, sector_id, board, listing_date,
                                   shares_outstanding, is_active, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)
                ON CONFLICT(ticker) DO UPDATE SET
                  yahoo_symbol       = excluded.yahoo_symbol,
                  name               = excluded.name,
                  sector_id          = excluded.sector_id,
                  board              = excluded.board,
                  listing_date       = excluded.listing_date,
                  shares_outstanding = excluded.shares_outstanding,
                  is_active          = 1,
                  updated_at         = excluded.updated_at
                """,
                (
                    r["ticker"],
                    f"{r['ticker']}.JK",
                    r["name"],
                    r["sector_id"],
                    r["board"],
                    r["listing_date"],
                    r["shares_outstanding"],
                    ts,
                ),
            )

        # Universe versi lama vs baru: pakai tabel temp supaya tak ada SQL yang dirakit string.
        conn.execute("CREATE TEMP TABLE _imported_ticker (ticker TEXT PRIMARY KEY)")
        conn.executemany(
            "INSERT OR IGNORE INTO _imported_ticker (ticker) VALUES (?)", ((r["ticker"],) for r in rows)
        )
        cur = conn.execute(
            """
            UPDATE stock SET is_active = 0, updated_at = ?
            WHERE is_active = 1 AND ticker NOT IN (SELECT ticker FROM _imported_ticker)
            """,
            (ts,),
        )
        deactivated = cur.rowcount

        if prev:
            import_id = prev["id"]
            conn.execute("DELETE FROM excel_import_error WHERE import_id = ?", (import_id,))
            conn.execute(
                "UPDATE excel_import SET filename = ?, imported_at = ?, sheet_count = ?, rows_ok = ?, rows_error = ? WHERE id = ?",
                (filename, ts, sheet_count, len(rows), len(errors), import_id),
            )
        else:
            cur = conn.execute(
                "INSERT INTO excel_import (filename, sha256, imported_at, sheet_count, rows_ok, rows_error) VALUES (?, ?, ?, ?, ?, ?)",
                (filename, sha, ts, sheet_count, len(rows), len(errors)),
            )
            import_id = cur.lastrowid

        if errors:
            conn.executemany(
                "INSERT INTO excel_import_error (import_id, sheet, row_no, raw_value, reason) VALUES (?, ?, ?, ?, ?)",
                [(import_id, e["sheet"], e["row_no"], e["raw_value"], e["reason"]) for e in errors],
            )
        conn.commit()
        return {
            "import_id": import_id,
            "sha256": sha,
            "sheet_count": sheet_count,
            "inserted": inserted,
            "updated": updated,
            "deactivated": deactivated,
            "errors": errors,
        }
    finally:
        conn.close()


def import_excel_path(path: str | Path, **kwargs) -> dict:
    p = Path(path)
    return import_excel(p.read_bytes(), filename=p.name, **kwargs)
